"""
Word 解析器
使用 python-docx 提取段落与表格文本
"""

from typing import List, Optional, Tuple

from src.document.parsers.base import Parser, ParsedContent


def _inline_image(doc, shape) -> Tuple[bytes, str]:
    """从 InlineShape 取出 (图片字节, 扩展名)

    为什么不用 `shape.image`：
        python-docx **1.2 的 InlineShape 没有 `.image` 属性**（`docx/shape.py`
        里只定义了 height/width/type）。旧写法 `getattr(shape, "image", None)`
        会静默兜到 None，于是 docx 内嵌图片其实**一张都没被提取过**——图片
        既不落附件、也进不了正文，而且完全没有任何报错。

    取件路径：
        wp:inline → a:graphic → a:graphicData → pic:pic → pic:blipFill
        → a:blip/@r:embed（关系 ID）→ doc.part.related_parts[rId]

    Raises:
        Exception: 非图片图形（图表/智能图形）取不到 blip 时抛出，由调用方跳过
    """
    blip = shape._inline.graphic.graphicData.pic.blipFill.blip
    r_id = blip.embed
    if not r_id:
        raise ValueError("图形无 r:embed 关系（可能是链接图片或非图片图形）")
    part = doc.part.related_parts[r_id]
    blob = part.blob
    name = str(getattr(part, "partname", "") or "")
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else "png"
    return blob, (ext or "png")


class DocxParser(Parser):
    """Word 解析器（.docx）"""

    extensions: List[str] = [".docx"]

    def parse_bytes(self, data: bytes, source_name: str) -> ParsedContent:
        import io

        from docx import Document as DocxDocument

        with io.BytesIO(data) as buf:
            doc = DocxDocument(buf)

        parts: List[str] = []

        # 段落
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)

        # 表格（逐行拼成文本行，便于检索）
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                line = " | ".join(c for c in cells if c)
                if line:
                    parts.append(line)

        content = "\n".join(parts)

        # 提取内嵌图片（多模态：附件存储，说明文字用固定标注）
        images = []
        for shape in doc.inline_shapes:
            try:
                blob, ext = _inline_image(doc, shape)
                if blob:
                    images.append({
                        "caption": "Word 文档内嵌图片",
                        "bytes": blob,
                        "ext": ext,
                    })
            except Exception:
                continue  # 单张图片提取失败（或非图片图形）不影响整体

        # 标题：首个非空段落（通常为首个标题/第一行）
        title: Optional[str] = None
        for para in doc.paragraphs:
            t = para.text.strip()
            if t:
                title = t if len(t) <= 80 else t[:80]
                break

        return ParsedContent(content=content, title=title, images=images)