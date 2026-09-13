"""
PowerPoint 解析器
使用 python-pptx 提取每页幻灯片中的文本框与表格文本
"""

from typing import List, Optional

from src.document.md_common import escape_pipe, render_md_table
from src.document.parsers.base import Parser, ParsedContent


class PptxParser(Parser):
    """PowerPoint 解析器（.pptx）"""

    extensions: List[str] = [".pptx"]

    def parse_bytes(self, data: bytes, source_name: str) -> ParsedContent:
        import io

        from pptx import Presentation

        with io.BytesIO(data) as buf:
            prs = Presentation(buf)

        slide_parts = []
        first_text: Optional[str] = None
        images = []

        for idx, slide in enumerate(prs.slides, 1):
            lines = []
            slide_text = []
            for shape in slide.shapes:
                # 文本框
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        lines.append(text)
                        slide_text.append(text)
                        if first_text is None and len(text) <= 100:
                            first_text = text
                # 表格
                if getattr(shape, "has_table", False):
                    # 与转换接口（to_markdown）共用同一渲染规则：补齐列宽 +
                    # 补 `| --- |` 分隔行。此前用 " | ".join(非空单元格) 拼成
                    # 普通文本行，列数不齐会错位、且 Obsidian 不认这是表格。
                    raw_rows = [
                        [escape_pipe(c.text.strip()) for c in row.cells]
                        for row in shape.table.rows
                    ]
                    table_md = render_md_table(raw_rows)
                    if table_md:
                        lines.append(table_md)
                        slide_text.append(table_md)
                # 图片（多模态：上下文字幕取本页文本）
                if str(getattr(shape, "shape_type", "")) == "PICTURE (13)" or "PICTURE" in str(getattr(shape, "shape_type", "")):
                    try:
                        images.append({
                            "caption": ("；".join(slide_text) or "（本页无文字说明）")[:200],
                            "bytes": shape.image.blob,
                            "ext": getattr(shape.image, "ext", "png"),
                        })
                    except Exception:
                        continue
            if lines:
                slide_parts.append(f"## 幻灯片 {idx}\n" + "\n".join(lines))

        content = "\n\n".join(slide_parts)

        # 标题：取首个短文本（通常为演示文稿标题页文字）
        title = first_text if first_text else (source_name if source_name else None)
        return ParsedContent(content=content, title=title, images=images)