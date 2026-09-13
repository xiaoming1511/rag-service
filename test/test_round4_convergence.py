"""
第四轮：document 层「收敛」回归

覆盖把重复实现收敛到 `src.document.md_common` 之后的行为：
    1. md_common.render_md_table / escape_pipe 本身
    2. 同一张逻辑表经 html / docx / pptx 三条路径 → 渲染结果一致
    3. 文本解码链统一（不再有三套写法）
    4. parsers/epub 改为结构感知（此前 get_text(" ") 会把 <pre> 压成一行）
    5. parsers/html 与 to_markdown 对同一份 HTML 产出同一份 Markdown

这些断言的作用是「锁死收敛结果」：任何一条路径被改回手写拼装，这里都会红。
"""

import io

import pytest


# ================================================================
# 1. md_common 单元
# ================================================================

class TestMdCommonTable:

    def test_header_and_separator_present(self):
        """必须有 `| --- |` 分隔行——缺了 Obsidian 不认这是表格"""
        from src.document.md_common import render_md_table

        md = render_md_table([["A", "B"], ["1", "2"]])
        assert md.splitlines()[0] == "| A | B |"
        assert md.splitlines()[1] == "| --- | --- |"
        assert md.splitlines()[2] == "| 1 | 2 |"

    def test_ragged_rows_padded_to_max_width(self):
        """列数不齐时右侧补空，避免列塌陷错位"""
        from src.document.md_common import render_md_table

        md = render_md_table([["A", "B", "C"], ["1"], ["x", "y"]])
        lines = md.splitlines()
        assert {ln.count("|") for ln in lines} == {4}, lines
        assert lines[2] == "| 1 |  |  |"
        assert lines[3] == "| x | y |  |"

    def test_blank_rows_dropped(self):
        from src.document.md_common import render_md_table

        md = render_md_table([["A", "B"], ["", ""], ["1", "2"]])
        assert len(md.splitlines()) == 3, md

    def test_empty_input_returns_empty(self):
        from src.document.md_common import render_md_table

        assert render_md_table([]) == ""
        assert render_md_table([[]]) == ""
        assert render_md_table([["", ""]]) == ""

    def test_escape_pipe_escapes_pipe_and_newline(self):
        from src.document.md_common import escape_pipe

        assert escape_pipe("a|b") == "a\\|b"
        assert escape_pipe("l1\nl2") == "l1 l2"
        assert escape_pipe("") == ""


# ================================================================
# 2. 三种格式的表格渲染一致性（收敛的核心断言）
# ================================================================

_CELLS = [["A1", "", "A3"], ["B1", "B2", ""]]
_EXPECTED_TABLE = (
    "| A1 |  | A3 |\n"
    "| --- | --- | --- |\n"
    "| B1 | B2 |  |"
)


def _html_table() -> str:
    rows = "".join(
        "<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in _CELLS
    )
    return f"<html><body><table>{rows}</table></body></html>"


def _docx_table() -> bytes:
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=len(_CELLS), cols=len(_CELLS[0]))
    for r, row in enumerate(_CELLS):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx_table() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    table = slide.shapes.add_table(len(_CELLS), len(_CELLS[0]), 0, 0, 400, 100).table
    for r, row in enumerate(_CELLS):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _extract_table(md: str) -> str:
    """抽出 Markdown 中的表格块（连续以 | 开头的行）"""
    out, in_table = [], False
    for ln in md.splitlines():
        if ln.startswith("|"):
            out.append(ln)
            in_table = True
        elif in_table:
            break
    return "\n".join(out)


class TestTableRenderingConvergence:

    def test_html_table_renders_expected(self):
        from src.document.to_markdown import convert_to_markdown

        md = convert_to_markdown(_html_table().encode("utf-8"), "html").markdown
        assert _extract_table(md) == _EXPECTED_TABLE, md

    def test_docx_table_renders_expected(self):
        from src.document.to_markdown import convert_to_markdown

        md = convert_to_markdown(_docx_table(), "docx").markdown
        assert _extract_table(md) == _EXPECTED_TABLE, md

    def test_pptx_table_renders_expected(self):
        """第四轮修正点：pptx 此前缺分隔行、且跳过空单元格导致列塌陷"""
        from src.document.to_markdown import convert_to_markdown

        md = convert_to_markdown(_pptx_table(), "pptx").markdown
        assert _extract_table(md) == _EXPECTED_TABLE, md

    def test_all_three_paths_agree(self):
        from src.document.to_markdown import convert_to_markdown

        html_md = _extract_table(convert_to_markdown(_html_table().encode("utf-8"), "html").markdown)
        docx_md = _extract_table(convert_to_markdown(_docx_table(), "docx").markdown)
        pptx_md = _extract_table(convert_to_markdown(_pptx_table(), "pptx").markdown)
        assert html_md == docx_md == pptx_md == _EXPECTED_TABLE


# ================================================================
# 3. 解码链统一
# ================================================================

class TestDecodeChain:

    def test_utf8_and_gb18030(self):
        from src.document.md_common import decode_bytes

        assert decode_bytes("简体中文测试".encode("utf-8")) == "简体中文测试"
        assert decode_bytes("简体中文测试".encode("gb18030")) == "简体中文测试"

    def test_never_raises_on_arbitrary_bytes(self):
        """链末 latin-1 兜底：任意字节都必须能解出（不抛异常）"""
        from src.document.md_common import decode_bytes

        for payload in (b"", b"\xff\xfe\x00", bytes(range(256))):
            assert isinstance(decode_bytes(payload), str)

    def test_all_text_entry_points_share_one_chain(self):
        """三条入口（html / to_markdown 文本类 / parsers 文本类）解码结果必须一致"""
        from src.document.html_to_markdown import HTMLToMarkdownConverter
        from src.document.parsers.text import TextParser
        from src.document.to_markdown import convert_to_markdown

        raw = "中文标题\n正文内容".encode("gb18030")

        via_convert = convert_to_markdown(raw, "txt", source_name="a.txt").markdown
        via_parser = TextParser().parse_bytes(raw, "a.txt").content
        assert "中文标题" in via_convert and "正文内容" in via_convert
        assert "中文标题" in via_parser and "正文内容" in via_parser

        html_raw = "<html><body><p>中文正文</p></body></html>".encode("gb18030")
        assert "中文正文" in HTMLToMarkdownConverter().convert_bytes(html_raw)

    def test_big5_limitation_is_explicit_not_silent(self):
        """已知局限：gb18030 会「成功」吃掉 Big5 字节，Big5 文本得到确定性乱码。

        这里锁的是「确定性」而非「正确性」——若将来引入字符集探测修好它，
        本用例会失败，提示把结论一并更新（而不是悄悄变了行为）。
        （R10-2：链中不可达的 big5 死分支已移除，行为与移除前完全一致。）
        """
        from src.document.md_common import decode_bytes

        correct = "繁體中文測試內容"
        got = decode_bytes(correct.encode("big5"))
        assert got != correct, "若此断言失败，说明已修好 Big5，请更新本用例与 md_common 注释"
        assert isinstance(got, str)  # 至少不是崩溃/异常


# ================================================================
# 4. EPUB 结构感知（索引路径）
# ================================================================

def _build_epub(chapter_html: str) -> bytes:
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("t-1")
    book.set_title("测试书")
    book.set_language("zh")
    c1 = epub.EpubHtml(title="章1", file_name="c1.xhtml", lang="zh")
    c1.content = chapter_html
    book.add_item(c1)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", c1]
    buf = io.BytesIO()
    epub.write_epub(buf, book)
    return buf.getvalue()


_CHAPTER = (
    "<html><body><h1>第一章</h1>"
    "<pre>┌───┐\n│ A │\n└───┘</pre>"
    "<p>正文段落</p></body></html>"
)


class TestEpubStructureAware:

    def test_pre_block_survives_as_code_fence(self):
        """此前 parsers/epub 用 get_text(" ") → <pre> 被压成一行、缩进全丢"""
        from src.document.parsers.epub import EPUBParser

        content = EPUBParser().parse_bytes(_build_epub(_CHAPTER), "t.epub").content
        assert "```" in content, f"EPUB 章节的 <pre> 没被保留成代码块:\n{content}"
        assert "┌───┐" in content
        # 换行必须还在（被压成一行是旧 bug 的特征）
        assert "│ A │" in content and content.index("┌───┐") < content.index("│ A │")

    def test_heading_survives(self):
        from src.document.parsers.epub import EPUBParser

        content = EPUBParser().parse_bytes(_build_epub(_CHAPTER), "t.epub").content
        assert "# 第一章" in content, content

    def test_title_from_dc_metadata(self):
        from src.document.parsers.epub import EPUBParser

        pc = EPUBParser().parse_bytes(_build_epub(_CHAPTER), "t.epub")
        assert pc.title == "测试书"

    def test_agrees_with_convert_api_on_code_fence(self):
        """索引路径与转换接口路径对同一本 EPUB 至少必须在「代码块保留」上一致"""
        from src.document.parsers.epub import EPUBParser
        from src.document.to_markdown import convert_to_markdown

        data = _build_epub(_CHAPTER)
        idx = EPUBParser().parse_bytes(data, "t.epub").content
        conv = convert_to_markdown(data, "epub", source_name="t.epub").markdown
        assert idx.count("```") == conv.count("```") == 2


# ================================================================
# 5. HTML 两条路径同源
# ================================================================

class TestHtmlPathParity:

    def test_parsers_html_and_to_markdown_agree(self):
        from src.document.parsers.html import HTMLParser
        from src.document.to_markdown import convert_to_markdown

        html = (
            "<html><head><title>标题X</title></head><body>"
            "<h1>标题X</h1><p>段落一</p>"
            "<pre>┌─┐\n│ │\n└─┘</pre>"
            "<table><tr><td>a</td><td>b</td></tr><tr><td>1</td><td>2</td></tr></table>"
            "</body></html>"
        ).encode("utf-8")

        parser_content = HTMLParser().parse_bytes(html, "x.html").content
        conv_md = convert_to_markdown(html, "html", source_name="x.html").markdown

        # 正文完全一致；转换接口额外把标题前缀成 H1（属既定差异，故用包含关系断言）
        assert parser_content.strip() in conv_md
        assert parser_content.count("```") == 2
        assert "| a | b |" in parser_content
