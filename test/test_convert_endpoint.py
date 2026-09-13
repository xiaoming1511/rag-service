"""
文档 → Markdown 统一转换端点测试（离线，无 pipeline / oMLX 依赖）

覆盖 /v1/convert/to-md：
- 各格式（html/txt/pdf/docx/pptx）转换 + 标题层级保留
- base64 二进制路径
- 格式校验 / 空内容守卫
- html2md 兼容端点 + formats 列表
"""

import base64
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import pytest
from fastapi.testclient import TestClient

from src.api.app import create_app


@pytest.fixture()
def client():
    return TestClient(create_app(None))


ARCH = """┌─────────────────────┐
│    接入层 (Interface) │
│  ┌─────────┐         │
│  │  Web UI │         │
│  └─────────┘         │
└─────────────────────┘
          │
          ▼
┌─────────────────────┐
│      API 层          │
└─────────────────────┘"""


def _make_docx() -> bytes:
    from docx import Document
    d = Document()
    d.add_heading("第一章 概述", 1)
    d.add_heading("1.1 背景", 2)
    d.add_paragraph("这是正文段落。")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _make_pptx() -> bytes:
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "演示标题"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _make_pdf() -> bytes:
    import pymupdf
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 72), "PDF Title", fontsize=24)
    p.insert_text((72, 110), "Body paragraph", fontsize=11)
    pdf = doc.tobytes()
    doc.close()
    return pdf


# ================================================================
# HTML（文本路径）
# ================================================================

def test_html_conversion(client):
    html = f'<html><head><title>T</title></head><body><h1>架构</h1><div class="architecture">{ARCH}</div></body></html>'
    r = client.post("/v1/convert/to-md", json={"format": "html", "content": html})
    assert r.status_code == 200
    d = r.json()
    assert d["ok"] is True
    assert "接入层" in d["markdown"]
    assert "API 层" in d["markdown"]
    assert d["markdown"].count("```") % 2 == 0


# ================================================================
# 二进制（base64 路径）
# ================================================================

def test_docx_conversion_preserves_heading_levels(client):
    data = _make_docx()
    b64 = base64.b64encode(data).decode()
    r = client.post("/v1/convert/to-md", json={
        "format": "docx", "content": b64, "content_is_base64": True,
    })
    assert r.status_code == 200
    md = r.json()["markdown"]
    assert "# 第一章 概述" in md
    assert "## 1.1 背景" in md
    assert "这是正文段落。" in md


def test_pptx_conversion(client):
    data = _make_pptx()
    b64 = base64.b64encode(data).decode()
    r = client.post("/v1/convert/to-md", json={
        "format": "pptx", "content": b64, "content_is_base64": True,
    })
    assert r.status_code == 200
    md = r.json()["markdown"]
    assert "幻灯片 1" in md
    assert "演示标题" in md


def test_pdf_conversion(client):
    data = _make_pdf()
    b64 = base64.b64encode(data).decode()
    r = client.post("/v1/convert/to-md", json={
        "format": "pdf", "content": b64, "content_is_base64": True,
    })
    assert r.status_code == 200
    md = r.json()["markdown"]
    assert "PDF Title" in md


def test_txt_conversion(client):
    r = client.post("/v1/convert/to-md", json={"format": "txt", "content": "标题行\n正文内容"})
    assert r.status_code == 200
    d = r.json()
    assert "标题行" in d["markdown"]
    assert d["title"] in (None, "标题行") or "标题行" in d["markdown"]


# ================================================================
# 守卫与兼容
# ================================================================

def test_unsupported_format_rejected(client):
    r = client.post("/v1/convert/to-md", json={"format": "xyz", "content": "x"})
    assert r.status_code == 400


def test_empty_content_rejected(client):
    r = client.post("/v1/convert/to-md", json={"format": "html", "content": "  "})
    assert r.status_code == 400


def test_bad_base64_rejected(client):
    r = client.post("/v1/convert/to-md", json={
        "format": "pdf", "content": "!!!not-base64!!!", "content_is_base64": True,
    })
    assert r.status_code == 400


def test_html2md_compat_endpoint(client):
    html = f'<html><body><h1>T</h1><div class="architecture">{ARCH}</div></body></html>'
    r = client.post("/v1/convert/html2md", json={"html": html})
    assert r.status_code == 200
    assert "接入层" in r.json()["markdown"]


def test_formats_list(client):
    r = client.get("/v1/convert/formats")
    assert r.status_code == 200
    fmts = r.json()["formats"]
    for f in ("html", "pdf", "docx", "pptx", "epub", "txt", "md"):
        assert f in fmts
