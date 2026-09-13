"""
OCR 集成回归（视觉模型 OvisOCR2）

覆盖：
  1. 配置层     —— 默认关闭、样例配置与 schema 同步、/v1/config 白名单
  2. 客户端护栏 —— 过小/超大/坏图跳过、格式归一化、失败降级、围栏剥离
  3. 索引路径   —— 扫描页整页 OCR、独立图片、内嵌图片（docx/pdf/epub）、各类上限
  4. 转换接口   —— 图片转 Markdown、扫描页回落、pptx/epub 内嵌图片
  5. 结构性保证 —— OCR 抛异常绝不影响文档入库

全部使用桩 OCR 客户端，不依赖真实模型服务；涉及 127.0.0.1:9 的用例只验证
「请求失败 → 返回空串」这条降级路径（该端口无监听，连接必然失败）。
"""

import base64
import io
from contextlib import contextmanager
from pathlib import Path

import pytest

from src.document.loader import DocumentLoader
from src.document.to_markdown import convert_to_markdown, supported_formats

ROOT = Path(__file__).resolve().parents[1]


# ================================================================
#  桩与素材构造
# ================================================================

class StubOCR:
    """桩 OCR 客户端（接口与 src.document.ocr.OCRClient 对齐）"""

    def __init__(self, text="桩识别出的文字", *, fail=False,
                 min_text_chars=16, max_pages_per_doc=30, max_images_per_doc=20,
                 render_dpi=150, prompt="OCR"):
        self.model = "stub-ocr"
        self.text = text
        self.fail = fail
        self.min_text_chars = min_text_chars
        self.max_pages_per_doc = max_pages_per_doc
        self.max_images_per_doc = max_images_per_doc
        self.render_dpi = render_dpi
        self.prompt = prompt
        self.page_calls = 0
        self.image_calls = 0

    def looks_scanned(self, page_text: str) -> bool:
        return len((page_text or "").strip()) < self.min_text_chars

    def ocr_pdf_page(self, page, *, source: str = "") -> str:
        self.page_calls += 1
        if self.fail:
            raise RuntimeError("桩故障")
        return self.text

    def ocr_image_bytes(self, data: bytes, *, ext: str = "png", source: str = "") -> str:
        self.image_calls += 1
        if self.fail:
            raise RuntimeError("桩故障")
        return self.text


def _img_bytes(fmt="PNG", size=(320, 200)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", size, "white").save(buf, format=fmt)
    return buf.getvalue()


def _png(size=(320, 200)) -> bytes:
    return _img_bytes("PNG", size)


def _write_lines(page, lines) -> None:
    """在页面上写入 ASCII 文本（PyMuPDF 内置字体不支持 CJK，故测试用英文）"""
    y = 72
    for ln in lines:
        page.insert_text((72, y), ln, fontsize=14)
        y += 28


def _text_pdf(lines=("Chapter 1 Configuration", "chunk size 800, overlap 100.")) -> bytes:
    """有文本层的普通 PDF"""
    import pymupdf

    doc = pymupdf.open()
    _write_lines(doc.new_page(), lines)
    data = doc.tobytes()
    doc.close()
    return data


def _scanned_pdf(pages=2) -> bytes:
    """纯图片页、无文本层（模拟扫描件）"""
    import pymupdf

    doc = pymupdf.open()
    for _ in range(pages):
        page = doc.new_page(width=400, height=300)
        page.insert_image(pymupdf.Rect(20, 20, 380, 280), stream=_png())
    data = doc.tobytes()
    doc.close()
    return data


def _mixed_pdf() -> bytes:
    """第 1 页有文本层、第 2 页是纯图片（正文 + 扫描附录）"""
    import pymupdf

    doc = pymupdf.open()
    _write_lines(doc.new_page(), ("Chapter 1 Configuration",))
    page2 = doc.new_page(width=400, height=300)
    page2.insert_image(pymupdf.Rect(20, 20, 380, 280), stream=_png())
    data = doc.tobytes()
    doc.close()
    return data


def _pdf_with_embedded_image() -> bytes:
    """文本页 + 一页上嵌了张图（文本层存在 → 图片应按内嵌图片走 OCR）"""
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    _write_lines(page, ("Chapter 2 Figures",))
    page.insert_image(pymupdf.Rect(72, 120, 372, 320), stream=_png())
    data = doc.tobytes()
    doc.close()
    return data


def _docx_with_images(png_paths, paragraphs=("Word 正文段落",)) -> bytes:
    from docx import Document as DocxDocument

    d = DocxDocument()
    for p in paragraphs:
        d.add_paragraph(p)
    for p in png_paths:
        d.add_picture(str(p))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pptx_with_image(png_path) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "演示文稿标题"
    slide.shapes.add_picture(str(png_path), Inches(1), Inches(2), Inches(3))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _epub_with_image(png_bytes) -> bytes:
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("ocr-test-book")
    book.set_title("测试电子书")
    book.set_language("zh")

    ch = epub.EpubHtml(title="第一章", file_name="chap_01.xhtml", lang="zh")
    ch.content = ("<html><body><h1>第一章</h1>"
                  "<p>这是 EPUB 电子书的正文内容。</p>"
                  "<img src='images/fig1.png'/></body></html>")
    book.add_item(ch)
    book.add_item(epub.EpubItem(uid="fig1", file_name="images/fig1.png",
                                media_type="image/png", content=png_bytes))
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", ch]

    buf = io.BytesIO()
    epub.write_epub(buf, book)
    return buf.getvalue()


def _loader(tmp_path, exts, client=None) -> DocumentLoader:
    """测试用加载器：附件目录固定到 tmp_path，避免污染仓库 data/attachments"""
    return DocumentLoader(
        source_dirs=[str(tmp_path)],
        extensions=exts,
        attachment_dir=str(tmp_path / "attachments"),
        ocr_client=client,
    )


@contextmanager
def _ocr_config(**overrides):
    """临时把全局配置的 ocr 段改成启用（退出时恢复原状）"""
    from src.config import AppConfig, config_manager

    saved = config_manager._config
    base = (saved or AppConfig()).model_dump()
    base["ocr"] = {**base.get("ocr", {}), "enabled": True, **overrides}
    config_manager._config = AppConfig(**base)
    try:
        yield config_manager.config
    finally:
        config_manager._config = saved


# ================================================================
#  1. 配置层
# ================================================================

class TestOCRConfig:

    def test_disabled_by_default(self):
        from src.config import AppConfig

        cfg = AppConfig()
        assert cfg.ocr.enabled is False
        assert cfg.ocr.model == "OvisOCR2"
        assert cfg.ocr.max_pages_per_doc >= 1
        assert cfg.ocr.max_images_per_doc >= 1

    def test_settings_yaml_matches_schema(self):
        """样例配置必须能被 schema 吃下（防止加了字段忘了同步 yaml）"""
        import yaml

        from src.config import AppConfig

        raw = yaml.safe_load((ROOT / "config" / "settings.yaml").read_text(encoding="utf-8"))
        assert "ocr" in raw, "settings.yaml 缺少 ocr 段"
        cfg = AppConfig(**raw)
        assert cfg.ocr.enabled is False, "样例配置必须默认关闭 OCR"
        assert cfg.ocr.model == "OvisOCR2"

    def test_image_extensions_not_enabled_by_default(self):
        """图片扩展名默认不在扫描清单里——避免「OCR 关着却反复扫描图片」"""
        import yaml

        raw = yaml.safe_load((ROOT / "config" / "settings.yaml").read_text(encoding="utf-8"))
        exts = {e.lower() for e in raw["documents"]["supported_extensions"]}
        assert exts.isdisjoint({".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff"})

    def test_get_client_none_when_disabled(self):
        from src.document.ocr import get_ocr_client

        assert get_ocr_client() is None

    def test_get_client_built_when_enabled(self):
        from src.document.ocr import OCRClient, get_ocr_client

        with _ocr_config():
            client = get_ocr_client()
        assert isinstance(client, OCRClient)
        assert client.model == "OvisOCR2"
        # base_url 为空时回落 omlx.base_url
        assert client.base_url.startswith("http")

    def test_get_client_after_config_change_is_fresh(self):
        """配置热改后必须拿到新客户端（客户端不缓存配置）"""
        from src.document.ocr import get_ocr_client

        with _ocr_config(model="OvisOCR2"):
            assert get_ocr_client().model == "OvisOCR2"
        with _ocr_config(model="AnotherOCR"):
            assert get_ocr_client().model == "AnotherOCR"

    def test_config_whitelist_allows_ocr(self):
        from src.api.routes.config import _ALLOWED_TOP_LEVEL

        assert "ocr" in _ALLOWED_TOP_LEVEL

    def test_public_config_exposes_ocr(self):
        from src.api.routes.config import _public_config

        with _ocr_config():
            assert _public_config()["ocr"]["enabled"] is True


# ================================================================
#  2. 客户端护栏
# ================================================================

class TestOCRClientGuards:

    def _client(self, **kw):
        from src.document.ocr import OCRClient

        kw.setdefault("model", "stub")
        kw.setdefault("base_url", "http://127.0.0.1:9")  # 无监听：只验证降级
        kw.setdefault("timeout", 1.0)
        return OCRClient(**kw)

    def test_looks_scanned(self):
        c = self._client(min_text_chars=16)
        assert c.looks_scanned("") is True
        assert c.looks_scanned("   ") is True
        assert c.looks_scanned("短") is True
        assert c.looks_scanned("这是一段足够长的正文内容，用来超过阈值。") is False

    def test_small_image_skipped(self):
        c = self._client(min_image_side=64)
        assert c._prepare(_png((32, 32)), "png") is None

    def test_oversize_image_skipped(self):
        c = self._client(max_image_pixels=1000)
        assert c._prepare(_png((320, 200)), "png") is None

    def test_broken_image_returns_none(self):
        c = self._client()
        assert c._prepare(b"this is not an image", "png") is None
        assert c._prepare(b"", "png") is None

    def test_png_passthrough_keeps_original_bytes(self):
        c = self._client(min_image_side=1, max_image_pixels=10 ** 9)
        raw = _png((320, 200))
        payload, mime, w, h = c._prepare(raw, "png")
        assert payload == raw, "png 应零重编码透传"
        assert mime == "image/png"
        assert (w, h) == (320, 200)

    def test_non_passthrough_format_converted_to_png(self):
        c = self._client(min_image_side=1, max_image_pixels=10 ** 9)
        for fmt, ext in (("GIF", "gif"), ("BMP", "bmp"), ("WEBP", "webp"), ("TIFF", "tiff")):
            payload, mime, w, h = c._prepare(_img_bytes(fmt, (320, 200)), ext)
            assert mime == "image/png", f"{fmt} 应转成 PNG"
            assert payload[:8] == b"\x89PNG\r\n\x1a\n", f"{fmt} 转换结果不是 PNG"
            assert (w, h) == (320, 200)

    def test_request_failure_degrades_to_empty(self):
        c = self._client()
        assert c.ocr_image_bytes(_png(), ext="png", source="t") == ""

    def test_strip_code_fence(self):
        from src.document.ocr import _strip_code_fence

        assert _strip_code_fence("```markdown\nabc\n```") == "abc"
        assert _strip_code_fence("```\nabc\n```") == "abc"
        assert _strip_code_fence("abc") == "abc"
        assert _strip_code_fence("") == ""


# ================================================================
#  3. 索引路径
# ================================================================

class TestScannedPdf:

    def test_disabled_scanned_pdf_still_skipped(self, tmp_path):
        """未启用 OCR：扫描件仍按「空文档」跳过（与接入前完全一致）"""
        (tmp_path / "scan.pdf").write_bytes(_scanned_pdf(1))
        assert _loader(tmp_path, [".pdf"]).load() == []

    def test_scanned_pages_are_ocrd(self, tmp_path):
        stub = StubOCR("这是扫描页识别出的内容")
        (tmp_path / "scan.pdf").write_bytes(_scanned_pdf(2))
        doc = _loader(tmp_path, [".pdf"], stub).load_single(str(tmp_path / "scan.pdf"))

        assert doc is not None
        assert "这是扫描页识别出的内容" in doc.content
        assert doc.metadata["ocr_pages"] == 2
        assert stub.page_calls == 2

    def test_scanned_page_embedded_image_not_ocrd_twice(self, tmp_path):
        """扫描页的内嵌图就是页面本身，不得再按内嵌图片 OCR 一次"""
        stub = StubOCR("页面文字")
        (tmp_path / "scan.pdf").write_bytes(_scanned_pdf(1))
        doc = _loader(tmp_path, [".pdf"], stub).load_single(str(tmp_path / "scan.pdf"))

        assert stub.page_calls == 1
        assert stub.image_calls == 0
        assert doc.content.count("页面文字") == 1

    def test_page_cap(self, tmp_path):
        stub = StubOCR("文字", max_pages_per_doc=1)
        (tmp_path / "scan.pdf").write_bytes(_scanned_pdf(3))
        doc = _loader(tmp_path, [".pdf"], stub).load_single(str(tmp_path / "scan.pdf"))

        assert doc.metadata["ocr_pages"] == 1
        assert stub.page_calls == 1

    def test_text_pdf_not_ocrd(self, tmp_path):
        stub = StubOCR()
        (tmp_path / "text.pdf").write_bytes(_text_pdf())
        doc = _loader(tmp_path, [".pdf"], stub).load_single(str(tmp_path / "text.pdf"))

        assert "Chapter 1 Configuration" in doc.content
        assert stub.page_calls == 0
        assert "ocr_pages" not in doc.metadata

    def test_mixed_pdf_only_scanned_page_ocrd(self, tmp_path):
        stub = StubOCR("扫描附录文字")
        (tmp_path / "mix.pdf").write_bytes(_mixed_pdf())
        doc = _loader(tmp_path, [".pdf"], stub).load_single(str(tmp_path / "mix.pdf"))

        assert "Chapter 1 Configuration" in doc.content
        assert "扫描附录文字" in doc.content
        assert stub.page_calls == 1

    def test_embedded_image_in_text_page_is_ocrd(self, tmp_path):
        stub = StubOCR("插图里的文字")
        (tmp_path / "fig.pdf").write_bytes(_pdf_with_embedded_image())
        doc = _loader(tmp_path, [".pdf"], stub).load_single(str(tmp_path / "fig.pdf"))

        assert stub.page_calls == 0, "有文本层的页面不该整页 OCR"
        assert stub.image_calls == 1
        assert "插图里的文字" in doc.content
        assert doc.metadata["ocr_images"] == 1


class TestStandaloneImage:

    def test_disabled_yields_no_document(self, tmp_path):
        (tmp_path / "shot.png").write_bytes(_png())
        assert _loader(tmp_path, [".png"]).load() == []

    def test_enabled_indexes_image_text(self, tmp_path):
        stub = StubOCR("截图里的文字内容")
        (tmp_path / "shot.png").write_bytes(_png())
        docs = _loader(tmp_path, [".png"], stub).load()

        assert len(docs) == 1
        assert docs[0].content == "截图里的文字内容"
        assert docs[0].metadata["title"] == "shot"
        assert docs[0].metadata["format"] == "png"
        assert stub.image_calls == 1

    def test_image_parser_registered_by_default(self):
        assert ".png" in DocumentLoader(source_dirs=[]).extensions
        assert ".webp" in DocumentLoader(source_dirs=[]).extensions


class TestEmbeddedImagesIndex:

    def test_docx_parser_extracts_images_without_ocr(self, tmp_path):
        """独立于 OCR 的真 bug 回归：python-docx 1.2 的 InlineShape 没有
        `.image` 属性，旧写法 getattr(shape, "image", None) 永远兜到 None，
        导致 docx 内嵌图片一张都没提取过（既不落附件也不进正文）。
        """
        png = tmp_path / "fig.png"
        png.write_bytes(_png())
        (tmp_path / "doc.docx").write_bytes(_docx_with_images([png]))

        doc = _loader(tmp_path, [".docx"]).load_single(str(tmp_path / "doc.docx"))

        assert len(doc.metadata["images"]) == 1
        saved = Path(doc.metadata["images"][0]["path"])
        assert saved.exists() and saved.stat().st_size > 0

    def test_docx_embedded_image_text_merged(self, tmp_path):
        png = tmp_path / "fig.png"
        png.write_bytes(_png())
        (tmp_path / "doc.docx").write_bytes(_docx_with_images([png]))

        stub = StubOCR("Word 图片里的文字")
        doc = _loader(tmp_path, [".docx"], stub).load_single(str(tmp_path / "doc.docx"))

        assert "Word 正文段落" in doc.content
        assert "Word 图片里的文字" in doc.content
        assert "Word 文档内嵌图片" in doc.content  # caption 一并进正文，便于溯源
        assert doc.metadata["ocr_images"] == 1
        assert stub.image_calls == 1

    def test_attachment_still_saved_when_ocr_on(self, tmp_path):
        """OCR 不改变图片落盘行为：附件照旧写入 metadata['images']"""
        png = tmp_path / "fig.png"
        png.write_bytes(_png())
        (tmp_path / "doc.docx").write_bytes(_docx_with_images([png]))

        doc = _loader(tmp_path, [".docx"], StubOCR("文字")).load_single(str(tmp_path / "doc.docx"))

        assert len(doc.metadata["images"]) == 1
        assert Path(doc.metadata["images"][0]["path"]).exists()

    def test_image_cap(self, tmp_path):
        pngs = []
        for i in range(3):
            p = tmp_path / f"fig{i}.png"
            p.write_bytes(_png())
            pngs.append(p)
        (tmp_path / "doc.docx").write_bytes(_docx_with_images(pngs))

        stub = StubOCR("文字", max_images_per_doc=1)
        doc = _loader(tmp_path, [".docx"], stub).load_single(str(tmp_path / "doc.docx"))

        assert doc.metadata["ocr_images"] == 1
        assert stub.image_calls == 1
        assert len(doc.metadata["images"]) == 3, "附件落盘不受 OCR 上限影响"

    def test_epub_parser_enumerates_images(self, tmp_path):
        """索引路径补齐 epub 图片枚举（与转换接口口径一致）"""
        from src.document.parsers.epub import EPUBParser

        parsed = EPUBParser().parse_bytes(_epub_with_image(_png()), "book.epub")
        assert len(parsed.images) == 1
        assert parsed.images[0]["ext"] == "png"
        assert parsed.images[0]["bytes"]

    def test_epub_image_ocr_in_index_path(self, tmp_path):
        (tmp_path / "book.epub").write_bytes(_epub_with_image(_png()))
        stub = StubOCR("电子书图片里的文字")
        doc = _loader(tmp_path, [".epub"], stub).load_single(str(tmp_path / "book.epub"))

        assert doc is not None
        assert "电子书图片里的文字" in doc.content
        assert doc.metadata["ocr_images"] == 1


# ================================================================
#  4. 转换接口
# ================================================================

class TestConvertOCR:

    def test_image_requires_ocr(self):
        with pytest.raises(ValueError, match="OCR"):
            convert_to_markdown(_png(), "png", "a.png")

    def test_image_converted_with_ocr(self):
        stub = StubOCR("# 识别标题\n\n正文内容")
        r = convert_to_markdown(_png(), "png", "扫描件截图.png", ocr_client=stub)

        assert r.source_format == "image"
        assert "识别标题" in r.markdown
        assert r.title == "扫描件截图"

    def test_route_format_string_yields_no_title(self):
        """经 HTTP 路由进来时 source_name 是格式串（"png"），不能当标题"""
        stub = StubOCR("文字")
        assert convert_to_markdown(_png(), "png", "png", ocr_client=stub).title is None

    def test_image_empty_ocr_raises(self):
        stub = StubOCR("")
        with pytest.raises(ValueError, match="未识别出文字"):
            convert_to_markdown(_png(), "png", "a.png", ocr_client=stub)

    def test_scanned_pdf_convert_uses_ocr(self):
        stub = StubOCR("扫描页文字")
        r = convert_to_markdown(_scanned_pdf(2), "pdf", "s.pdf", ocr_client=stub)

        assert "扫描页文字" in r.markdown
        assert stub.page_calls == 2
        assert r.title == "扫描页文字"  # 无字号信息时用首段 OCR 文字兜底

    def test_text_pdf_convert_unchanged(self):
        stub = StubOCR()
        r = convert_to_markdown(_text_pdf(), "pdf", "t.pdf", ocr_client=stub)

        assert "Chapter 1 Configuration" in r.markdown
        assert stub.page_calls == 0

    def test_pptx_embedded_image_ocr_inline(self, tmp_path):
        png = tmp_path / "fig.png"
        png.write_bytes(_png())
        stub = StubOCR("幻灯片图片文字")
        r = convert_to_markdown(_pptx_with_image(png), "pptx", "p.pptx", ocr_client=stub)

        assert "幻灯片图片文字" in r.markdown
        assert stub.image_calls == 1
        # OCR 文字应落在同一页幻灯片段落内（紧随该页文本框）
        slide_part = r.markdown.split("## 幻灯片")[1]
        assert "幻灯片图片文字" in slide_part
        # 响应载荷仍是 base64，未混入原始 bytes（否则无法 JSON 序列化）
        assert r.images and r.images[0]["data_base64"]
        assert "bytes" not in r.images[0]

    def test_epub_embedded_image_ocr(self):
        stub = StubOCR("电子书图片文字")
        r = convert_to_markdown(_epub_with_image(_png()), "epub", "b.epub", ocr_client=stub)

        assert "电子书图片文字" in r.markdown
        assert stub.image_calls == 1
        assert "bytes" not in r.images[0]

    def test_disabled_convert_paths_unaffected(self):
        """未启用 OCR 时，非图片格式的转换行为不变"""
        r = convert_to_markdown(b"# Title\n\nbody", "md", "a.md")
        assert r.markdown.startswith("# Title")

    def test_supported_formats_include_images(self):
        fmts = supported_formats()
        for ext in ("png", "jpg", "jpeg", "webp", "gif", "bmp", "tiff", "tif"):
            assert ext in fmts


class TestConvertOCRRoute:
    """/v1/convert 路由层：错误映射与 ocr 段可见性"""

    @pytest.fixture()
    def client(self):
        from fastapi.testclient import TestClient

        from src.api.app import create_app

        return TestClient(create_app(None))

    def test_image_without_ocr_returns_400_with_reason(self, client):
        """图片格式未启用 OCR 时给明确 400，而不是笼统的 422"""
        resp = client.post("/v1/convert/to-md", json={
            "format": "png",
            "content": base64.b64encode(_png()).decode(),
            "content_is_base64": True,
        })

        assert resp.status_code == 400
        assert "OCR" in resp.json()["detail"]

    def test_convert_formats_lists_images(self, client):
        resp = client.get("/v1/convert/formats")
        assert resp.status_code == 200
        assert "png" in resp.json()["formats"]


@pytest.fixture()
def client_for_config():
    from fastapi.testclient import TestClient

    from src.api.app import create_app

    return TestClient(create_app(None))


class TestOCRConfigRoute:
    """/v1/config 对 ocr 段的读与写"""

    def test_get_config_exposes_ocr(self, client_for_config):
        with _ocr_config():
            data = client_for_config.get("/v1/config").json()
        assert data["ocr"]["enabled"] is True
        assert data["ocr"]["model"] == "OvisOCR2"

    def test_post_config_accepts_ocr_and_persists(self, client_for_config, tmp_path):
        from src.api.routes.config import set_pipeline
        from src.config import AppConfig, config_manager

        saved_cfg = config_manager._config
        saved_path = config_manager._config_path
        cfg_file = tmp_path / "settings.yaml"
        cfg_file.write_text("{}", encoding="utf-8")
        base = (saved_cfg or AppConfig()).model_dump()
        config_manager._config = AppConfig(**base)
        config_manager._config_path = cfg_file
        # POST /v1/config 要求 pipeline 已注入；本次只改 ocr 段，_apply_hot 不会
        # 触碰任何运行组件，因此一个空对象就足够（不需要真的建 pipeline）
        set_pipeline(object())
        try:
            resp = client_for_config.post("/v1/config", json={
                "ocr": {"enabled": True, "model": "OvisOCR2", "max_pages_per_doc": 7},
            })

            assert resp.status_code == 200, resp.text
            assert config_manager.config.ocr.enabled is True
            assert config_manager.config.ocr.max_pages_per_doc == 7

            import yaml

            saved = yaml.safe_load(cfg_file.read_text(encoding="utf-8"))
            assert saved["ocr"]["enabled"] is True, "ocr 段必须落盘"
            assert saved["ocr"]["max_pages_per_doc"] == 7
        finally:
            set_pipeline(None)
            config_manager._config = saved_cfg
            config_manager._config_path = saved_path


# ================================================================
#  5. 结构性保证：OCR 异常不影响入库
# ================================================================

class TestFailureTolerance:

    def test_embedded_image_exception_keeps_document(self, tmp_path):
        png = tmp_path / "fig.png"
        png.write_bytes(_png())
        (tmp_path / "doc.docx").write_bytes(_docx_with_images([png]))

        doc = _loader(tmp_path, [".docx"], StubOCR(fail=True)).load_single(str(tmp_path / "doc.docx"))

        assert doc is not None, "OCR 抛错不得让整篇文档入库失败"
        assert "Word 正文段落" in doc.content
        assert "ocr_images" not in doc.metadata
        assert len(doc.metadata["images"]) == 1  # 附件照旧落盘

    def test_scanned_page_exception_keeps_text_pages(self, tmp_path):
        (tmp_path / "mix.pdf").write_bytes(_mixed_pdf())
        doc = _loader(tmp_path, [".pdf"], StubOCR(fail=True)).load_single(str(tmp_path / "mix.pdf"))

        assert doc is not None
        assert "Chapter 1 Configuration" in doc.content
        # 失败页不计入"成功页数"，但必须留下可被运维看见的痕迹
        assert "ocr_pages" not in doc.metadata
        assert doc.metadata["ocr_pages_failed"] == 1

    def test_standalone_image_exception_skips_only_that_file(self, tmp_path):
        (tmp_path / "a.png").write_bytes(_png())
        (tmp_path / "ok.md").write_text("# 正常文档\n\n内容", encoding="utf-8")

        docs = _loader(tmp_path, [".png", ".md"], StubOCR(fail=True)).load()

        assert {d.file_name for d in docs} == {"ok.md"}

    def test_output_has_no_raw_bytes_leak_in_index_metadata(self, tmp_path):
        """metadata 里只应有附件路径，不应把图片原始字节带进向量库元数据"""
        png = tmp_path / "fig.png"
        png.write_bytes(_png())
        (tmp_path / "doc.docx").write_bytes(_docx_with_images([png]))

        doc = _loader(tmp_path, [".docx"], StubOCR("文字")).load_single(str(tmp_path / "doc.docx"))

        for entry in doc.metadata["images"]:
            assert "bytes" not in entry
