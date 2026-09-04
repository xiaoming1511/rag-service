"""
多格式解析测试（离线：自造样例文件 + 本地 HTTP 服务，无需 oMLX）

覆盖：
- Markdown（含 frontmatter）/ 纯文本 / PDF / Word / HTML 的加载解析
- 扩展名过滤
- URL 抓取（本地 HTTP 服务）与网页索引（假嵌入器）
"""

import sys
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

# 添加项目根目录到 Python 路径
sys.path.insert(0, str(Path(__file__).parent.parent))

import pytest

from src.document.loader import DocumentLoader
from src.document.chunker import Chunker
from src.vector_store.chroma_store import ChromaStore
from src.pipeline.indexer import Indexer


class FakeEmbedder:
    """离线假嵌入器：固定向量，避免网络依赖"""

    def embed(self, texts):
        return [[0.1] * 1024 for _ in texts]

    def embed_single(self, text):
        return [0.1] * 1024

    def get_embedding_dimension(self):
        return 1024


def _make_pdf(path: Path, lines):
    """用 PyMuPDF 生成一个简单 PDF"""
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "\n".join(lines), fontsize=12)
    # 中文字体需嵌入字体才能直接插入，这里用拉丁文字避免依赖系统字体
    doc.save(str(path))
    doc.close()


def _make_docx(path: Path, title: str, body: str):
    """用 python-docx 生成一个简单 .docx"""
    from docx import Document as DocxDocument
    doc = DocxDocument()
    doc.add_heading(title, level=1)
    doc.add_paragraph(body)
    doc.save(str(path))


@pytest.fixture()
def sample_vault(tmp_path) -> Path:
    """构建包含各格式样例的临时文档库"""
    src = tmp_path / "kb"
    src.mkdir()

    (src / "note.md").write_text(
        "---\ntitle: Markdown 示例\n---\n# 标题\n\n这是 Markdown 内容。",
        encoding="utf-8",
    )
    (src / "plain.txt").write_text(
        "纯文本标题\n这是纯文本的内容，用于验证 TXT 解析。",
        encoding="utf-8",
    )
    (src / "meeting.docx")  # 占位，下面生成
    _make_docx(src / "meeting.docx", "会议纪要", "这是 Word 文档的正文内容。")
    _make_pdf(src / "report.pdf", ["PDF Report", "This is PDF content for testing."])
    (src / "page.html").write_text(
        "<html><head><title>网页示例</title></head><body>"
        "<h1>网页标题</h1><p>这是 HTML 网页的正文段落。</p>"
        "<p>第二个段落。</p><script>var x = 1;</script></body></html>",
        encoding="utf-8",
    )
    return src


def test_load_all_formats(sample_vault):
    """加载器应解析出全部 5 种格式"""
    loader = DocumentLoader(source_dirs=[str(sample_vault)])
    docs = {d.file_name: d for d in loader.load()}

    # 5 个文档（md/txt/docx/pdf/html）
    assert set(docs) == {"note.md", "plain.txt", "meeting.docx", "report.pdf", "page.html"}

    # 各格式内容提取
    assert "这是 Markdown 内容" in docs["note.md"].content
    assert "这是纯文本的内容" in docs["plain.txt"].content
    assert "这是 Word 文档的正文内容" in docs["meeting.docx"].content
    assert "This is PDF content for testing" in docs["report.pdf"].content
    assert "这是 HTML 网页的正文段落" in docs["page.html"].content

    # 元数据：格式标记与标题
    assert docs["note.md"].metadata.get("format") == "md"
    assert docs["meeting.docx"].metadata.get("format") == "docx"
    assert docs["note.md"].metadata.get("title") == "Markdown 示例"  # frontmatter 优先
    assert docs["page.html"].metadata.get("title") == "网页示例"     # <title> 提取
    assert docs["plain.txt"].metadata.get("title") == "纯文本标题"   # 首行提取

    # HTML 噪声（script）应被剔除
    assert "var x = 1" not in docs["page.html"].content


def test_extensions_filter(sample_vault):
    """扩展名过滤：只加载指定格式"""
    loader = DocumentLoader(source_dirs=[str(sample_vault)], extensions=[".md", ".txt"])
    names = {d.file_name for d in loader.load()}
    assert names == {"note.md", "plain.txt"}


def test_unsupported_extension_ignored(tmp_path):
    """无解析器的扩展名（如 .xyz）应被忽略"""
    src = tmp_path / "kb"
    src.mkdir()
    (src / "data.xyz").write_text("未知格式内容", encoding="utf-8")
    loader = DocumentLoader(source_dirs=[str(src)])
    assert loader.load() == []


# ================================================================
# URL 抓取测试（本地 HTTP 服务）
# ================================================================

class _HtmlHandler(BaseHTTPRequestHandler):
    """最小 HTTP 服务：返回固定 HTML"""

    def do_GET(self):
        body = (
            "<html><head><title>本地测试页</title></head><body>"
            "<h1>本地标题</h1><p>这是本地 HTTP 服务返回的网页内容。</p>"
            "</body></html>"
        ).encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *args):
        pass  # 静默日志


@pytest.fixture()
def http_server():
    """启动本地 HTTP 服务，返回 base_url"""
    server = HTTPServer(("127.0.0.1", 0), _HtmlHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base = f"http://127.0.0.1:{server.server_address[1]}"
    yield base
    server.shutdown()
    thread.join(timeout=3)


def test_load_url(http_server):
    """URL 抓取解析：内容、标题、来源元数据"""
    loader = DocumentLoader(source_dirs=[])
    doc = loader.load_url(f"{http_server}/page")

    assert doc is not None
    assert "这是本地 HTTP 服务返回的网页内容" in doc.content
    assert doc.metadata.get("title") == "本地测试页"
    assert doc.metadata.get("source") == "url"
    assert doc.file_path.startswith("http://")
    assert loader.doc_id_for(doc.file_path) == doc.id  # URL 文档 ID 稳定


def test_index_url(http_server, tmp_path):
    """网页索引：分块入库，URL 文档 ID 可检索"""
    loader = DocumentLoader(source_dirs=[])
    chunker = Chunker(chunk_size=800, overlap=100, strategy="heading")
    store = ChromaStore(
        collection_name="test_url",
        persist_directory=str(tmp_path / "chroma_db"),
        embedding_dimension=1024,
    )
    indexer = Indexer(
        loader=loader,
        chunker=chunker,
        embedder=FakeEmbedder(),
        vector_store=store,
    )

    result = indexer.index_url(f"{http_server}/page")
    assert result.get("success") is True
    assert result.get("chunks", 0) >= 1
    assert store.count() == result["chunks"]

    # URL 文档 ID 可检索到其块
    doc_id = loader.doc_id_for(f"{http_server}/page")
    blocks = store.get_by_doc_id(doc_id)
    assert len(blocks) == result["chunks"]

    # 重复索引应跳过
    result2 = indexer.index_url(f"{http_server}/page")
    assert result2.get("skipped") is True


# ================================================================
# EPUB / PPTX 解析（自造样例）
# ================================================================

def _make_epub(path: Path, title: str = "测试电子书"):
    """用 ebooklib 生成最小 EPUB"""
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("test-epub-001")
    book.set_title(title)
    book.set_language("zh")
    book.add_author("测试作者")

    chap = epub.EpubHtml(title="第一章", file_name="chap1.xhtml", lang="zh")
    chap.content = "<h1>第一章</h1><p>这是 EPUB 电子书的正文内容，用于验证解析。</p>"
    book.add_item(chap)
    book.toc = (chap,)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", chap]

    epub.write_epub(str(path), book)


def _make_pptx(path: Path):
    """用 python-pptx 生成最小 PPTX"""
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "演示文稿标题"
    slide.placeholders[1].text = "这是 PPT 的正文要点。"

    table_slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = table_slide.shapes.add_table(2, 2, 72, 72, 200, 80).table
    table.cell(0, 0).text = "主题"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "性能"
    table.cell(1, 1).text = "优化后提升明显"

    prs.save(str(path))


def test_load_epub(sample_vault):
    """EPUB 解析：章节正文与标题元数据"""
    epub_path = sample_vault.parent / "book.epub"
    _make_epub(epub_path)

    loader = DocumentLoader(source_dirs=[str(sample_vault.parent)])
    docs = {d.file_name: d for d in loader.load()}

    book = docs.get("book.epub")
    assert book is not None
    assert "这是 EPUB 电子书的正文内容" in book.content
    assert book.metadata.get("title") == "测试电子书"
    assert book.metadata.get("format") == "epub"


def test_load_pptx(sample_vault):
    """PPTX 解析：文本框与表格内容"""
    pptx_path = sample_vault.parent / "slides.pptx"
    _make_pptx(pptx_path)

    loader = DocumentLoader(source_dirs=[str(sample_vault.parent)])
    docs = {d.file_name: d for d in loader.load()}

    slides = docs.get("slides.pptx")
    assert slides is not None
    assert "演示文稿标题" in slides.content
    assert "这是 PPT 的正文要点" in slides.content
    assert "优化后提升明显" in slides.content  # 表格内容
    assert slides.metadata.get("format") == "pptx"