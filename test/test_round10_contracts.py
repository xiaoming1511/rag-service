"""
Round 10 契约测试（三项低危记录项收尾）

- R10-1 history 轮对齐裁剪：token 裁剪不再裁出「以 assistant 开头」的轮中孤儿
- R10-2 解码链死分支移除：big5 从 _DECODE_CHAIN 移除（行为不变，代码诚实）
- R10-3 转换接口图片元数据回传：pptx 图片不再静默丢弃，epub 图片可枚举
"""

import base64
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

# 标准 1x1 PNG
PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


# ================================================================
# R10-1 history 轮对齐裁剪
# ================================================================

class TestHistoryTurnAlignedTrim:
    def _hist(self):
        return [
            {"role": "user", "content": "u" * 100},       # 50 tok
            {"role": "assistant", "content": "a" * 80},   # 40 tok
            {"role": "user", "content": "第二问"},         # 5 tok
            {"role": "assistant", "content": "第二答"},    # 5 tok
        ]

    def test_no_orphan_assistant_head(self):
        """预算只够裁 1 条时，必须连着同轮的 assistant 一起裁（旧实现会留下孤儿）"""
        from src.generation.history import trim_history

        out = trim_history(self._hist(), max_rounds=10, token_budget=91)
        assert [m["role"] for m in out] == ["user", "assistant"], \
            f"应按轮对齐裁剪，实际: {[m['role'] for m in out]}"
        assert out[0]["content"] == "第二问"

    def test_last_turn_always_intact(self):
        """无论如何裁剪，最近一问一答必须完整保留"""
        from src.generation.history import trim_history

        hist = self._hist()
        out = trim_history(hist, max_rounds=10, token_budget=1)
        assert len(out) == 2
        assert out[-1]["content"] == "第二答" and out[-2]["content"] == "第二问"

    def test_odd_length_degrades_gracefully(self):
        """奇数长度的畸形历史：保底最近 1 轮（2 条），不崩溃不丢最后一问"""
        from src.generation.history import trim_history

        hist = [
            {"role": "user", "content": "x" * 200},
            {"role": "assistant", "content": "y" * 200},
            {"role": "user", "content": "最后一问"},
        ]
        out = trim_history(hist, max_rounds=10, token_budget=10)
        assert out[-1]["content"] == "最后一问"
        assert len(out) >= 1

    def test_input_not_mutated(self):
        from src.generation.history import trim_history

        hist = self._hist()
        snapshot = [dict(m) for m in hist]
        trim_history(hist, max_rounds=10, token_budget=1)
        assert hist == snapshot

    def test_within_budget_untouched(self):
        from src.generation.history import trim_history

        out = trim_history(self._hist(), max_rounds=10, token_budget=10000)
        assert len(out) == 4


# ================================================================
# R10-2 解码链死分支移除
# ================================================================

class TestDecodeChainDeadBranch:
    def test_big5_removed_from_chain(self):
        """big5 在 gb18030 之后不可达（gb18030 字节空间覆盖任意双字节序列），
        留在链里只會误导读者以为它能救繁体文本——移除死分支，行为不变"""
        from src.document import md_common

        assert "big5" not in md_common._DECODE_CHAIN

    def test_behavior_unchanged_deterministic(self):
        """移除后行为必须与此前完全一致：Big5 文本 → 确定性 gb18030 乱码"""
        from src.document.md_common import decode_bytes

        correct = "繁體中文測試內容"
        got1 = decode_bytes(correct.encode("big5"))
        got2 = decode_bytes(correct.encode("big5"))
        assert got1 != correct, "若此断言失败，说明 Big5 已被正确解码，请更新 round4 锁定用例"
        assert got1 == got2, "解码必须是确定性的"

    def test_chain_still_covers_utf8_gb18030_latin1(self):
        from src.document.md_common import decode_bytes, _DECODE_CHAIN

        assert _DECODE_CHAIN[0] == "utf-8" and _DECODE_CHAIN[-1] == "latin-1"
        assert "gb18030" in _DECODE_CHAIN
        assert decode_bytes("简体中文".encode("gb18030")) == "简体中文"


# ================================================================
# R10-3 转换接口图片元数据回传
# ================================================================

def _make_pptx_with_picture() -> bytes:
    """最小 PPTX：第一页有说明文字 + 一张图片"""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式
    tb = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(914400), Emu(914400))
    tb.text_frame.text = "本页图片说明文字"
    slide.shapes.add_picture(io.BytesIO(PNG_1PX), Emu(914400), Emu(0))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _make_epub_with_image() -> bytes:
    """最小 EPUB：一章正文 + 一张图片资源"""
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("r10-epub")
    book.set_title("R10 测试书")
    book.set_language("zh")
    chap = epub.EpubHtml(title="第一章", file_name="chap1.xhtml", lang="zh")
    chap.content = '<h1>第一章</h1><p>正文</p><img src="static/cover.png"/>'
    book.add_item(chap)
    img = epub.EpubImage()
    img.file_name = "static/cover.png"
    img.content = PNG_1PX
    book.add_item(img)
    book.toc = (chap,)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", chap]
    buf = io.BytesIO()
    epub.write_epub(buf, book)
    return buf.getvalue()


class TestConvertImagesMetadata:
    def test_pptx_pictures_extracted(self):
        """pptx 图片不再静默丢弃：ConvertResult.images 携带说明/格式/base64"""
        from src.document.to_markdown import convert_to_markdown

        result = convert_to_markdown(_make_pptx_with_picture(), "pptx", source_name="a.pptx")
        assert len(result.images) == 1, "pptx 中的一张图片必须被回传"
        img = result.images[0]
        assert img["slide"] == 1
        assert img["ext"] == "png"
        assert "本页图片说明文字" in img["caption"]
        assert base64.b64decode(img["data_base64"]) == PNG_1PX

    def test_epub_images_enumerated(self):
        """epub 图片资源可枚举：src 与正文引用对应，data 可取出"""
        from src.document.to_markdown import convert_to_markdown

        result = convert_to_markdown(_make_epub_with_image(), "epub", source_name="b.epub")
        assert len(result.images) == 1
        img = result.images[0]
        assert img["src"] == "static/cover.png"
        assert img["ext"] == "png"
        assert base64.b64decode(img["data_base64"]) == PNG_1PX

    def test_api_response_carries_images(self):
        """/v1/convert/to-md 响应回传 images 字段（加法式，旧消费方无感）"""
        import base64 as b64
        from fastapi.testclient import TestClient

        from src.api.app import create_app

        client = TestClient(create_app(None))
        payload = b64.b64encode(_make_pptx_with_picture()).decode()
        resp = client.post("/v1/convert/to-md", json={
            "format": "pptx", "content": payload, "content_is_base64": True,
        })
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert len(body["images"]) == 1
        assert body["images"][0]["ext"] == "png"
        assert body["stats"]["images_count"] == 1

    def test_payload_guard_drops_data_keeps_metadata(self, monkeypatch):
        """图片载荷超聚合上限：去 data 保元数据（响应不膨胀，信息不丢失）"""
        from src.api.routes import convert as convert_route

        monkeypatch.setattr(convert_route, "_MAX_IMAGES_PAYLOAD", 150)
        images = [
            {"slide": 1, "caption": "c1", "ext": "png", "size_bytes": 100,
             "data_base64": "A" * 100},
            {"slide": 2, "caption": "c2", "ext": "png", "size_bytes": 100,
             "data_base64": "B" * 100},
        ]
        shaped, truncated = convert_route._shape_images(images)
        assert truncated is True
        assert shaped[0]["data_base64"] == "A" * 100, "预算内的图片必须保留数据"
        assert shaped[1]["data_base64"] is None, "超限图片只保留元数据"
        assert shaped[1]["caption"] == "c2"
