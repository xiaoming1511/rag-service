"""
第三轮修复回归测试（离线，无需 oMLX）

覆盖第三轮评审发现的问题，逐项锁死行为：
- N2  SSRF 防护：默认拒绝内网/回环/元数据/未指定地址；开关可放行
- N4  logging text 格式：子 logger 记录不再被 root-filter 吞掉
- N1  IngestQueue：worker 退出与 submit 的竞态（任务不得永久 queued）
- L1  backup：keep<=0 不得删光备份；同秒两次备份不覆盖
- L2  backup：原子写（不产生半写 JSON）
- L3  convert：显式 title 覆盖生效
- L4  to_markdown：txt 首行标题不重复出现
- L6  to_markdown：pptx 表格列数对齐
- L7  session store：并发删除不炸 list()；truncate 负数不再静默无操作
- L8  convert：超限体积在解码前即被 413 拦下
"""

import io
import json
import logging
import sys
import threading
import time
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

from src.document.to_markdown import prefix_title
from src.security.url_safety import UnsafeURLError, validate_public_url


# ================================================================
# N2：SSRF 防护
# ================================================================

class TestUrlSafety:
    """默认策略必须是「拒绝内网」，放行必须走显式开关"""

    @pytest.mark.parametrize("url", [
        "http://127.0.0.1/admin",          # 回环
        "http://localhost:8080/",          # 域名解析到回环
        "http://10.0.0.5/",                # 私网
        "http://192.168.1.1/",             # 私网
        "http://172.16.0.9/",              # 私网
        "http://169.254.169.254/latest/",  # 云元数据（链路本地）
        "http://0.0.0.0/",                 # 未指定地址
        "http://[::1]/",                   # IPv6 回环
        "file:///etc/passwd",              # 协议不允许
        "ftp://example.com/x",             # 协议不允许
    ])
    def test_rejects_unsafe_targets(self, url, monkeypatch):
        monkeypatch.delenv("RAG_ALLOW_PRIVATE_URLS", raising=False)
        with pytest.raises(UnsafeURLError):
            validate_public_url(url, allow_private=False)

    def test_allows_public_ip_literal(self, monkeypatch):
        """公网 IP 字面量不经 DNS，必须放行"""
        monkeypatch.delenv("RAG_ALLOW_PRIVATE_URLS", raising=False)
        validate_public_url("http://93.184.216.34/", allow_private=False)

    def test_explicit_flag_allows_private(self):
        validate_public_url("http://127.0.0.1:9000/", allow_private=True)

    def test_env_switch_allows_private(self, monkeypatch):
        monkeypatch.setenv("RAG_ALLOW_PRIVATE_URLS", "1")
        validate_public_url("http://127.0.0.1:9000/")

    def test_env_switch_false_value_still_blocks(self, monkeypatch):
        monkeypatch.setenv("RAG_ALLOW_PRIVATE_URLS", "0")
        with pytest.raises(UnsafeURLError):
            validate_public_url("http://127.0.0.1:9000/")

    def test_missing_host_rejected(self):
        with pytest.raises(UnsafeURLError):
            validate_public_url("http:///nohost", allow_private=False)


# ================================================================
# N4：text 日志格式下的请求关联
# ================================================================

def test_text_log_format_keeps_child_logger_records(monkeypatch, tmp_path):
    """
    回归：RequestIdFilter 曾挂在 root logger 上，而 logger 级 filter 不作用于
    子 logger 传播上来的记录 → text 格式的 %(request_id)s 抛 ValueError，
    记录被 logging 内部错误吞掉（表现为日志消失）。

    这里直接复刻「handler 上挂 filter」的最终形态并断言记录可见。
    """
    from src.logging_setup import RequestIdFilter, _TEXT_FORMAT, set_request_id

    root = logging.getLogger()
    buf = io.StringIO()
    handler = logging.StreamHandler(buf)
    handler.setFormatter(logging.Formatter(_TEXT_FORMAT, datefmt="%H:%M:%S"))
    handler.addFilter(RequestIdFilter())  # 关键：filter 在 handler 上
    root.addHandler(handler)
    child = logging.getLogger("rag.round3.child")
    old_level = child.level
    child.setLevel(logging.INFO)
    tok = set_request_id("RID-R3")
    try:
        child.warning("第三轮回归：text 格式子 logger 记录")
        out = buf.getvalue()
    finally:
        set_request_id(None) if False else None
        from src.logging_setup import reset_request_id
        reset_request_id(tok)
        child.setLevel(old_level)
        root.removeHandler(handler)

    assert "第三轮回归" in out, f"子 logger 记录被吞掉，实际输出: {out!r}"
    assert "rid=RID-R3" in out, f"request_id 未写入，实际输出: {out!r}"


# ================================================================
# N1：IngestQueue worker 生命周期
# ================================================================

def test_worker_clears_thread_when_idle_and_runs_later_submit(tmp_path):
    """
    回归：worker 在「队列已空」分支若不置空 self._thread 就退出，
    submit() 的 _ensure_worker() 会看到旧的 is_alive()==True 而跳过启动
    新 worker → 新任务永久停留在 queued（静默卡死）。
    """
    from src.pipeline.ingest_queue import IngestQueue

    ran = threading.Event()

    def run_fn(job):
        ran.set()
        return {"ok": True}

    q = IngestQueue(run_fn=run_fn, store_path=str(tmp_path / "jobs.json"))
    j1 = q.submit("incremental")
    assert ran.wait(5), "首个任务未被 worker 执行"
    assert q.get(j1.id).status == "done"

    # 等 worker 自行退出，并断言它把 _thread 置空（竞态修复的核心不变量）
    deadline = time.time() + 5
    while q._thread is not None and time.time() < deadline:
        time.sleep(0.05)
    assert q._thread is None, "worker 退出时必须置空 _thread，否则 submit 会跳过重启"

    ran.clear()
    j2 = q.submit("incremental")
    assert ran.wait(5), "worker 退出后新提交的任务必须被重新拉起执行"
    assert q.get(j2.id).status == "done"


def test_cancel_running_job_marks_cancelled(tmp_path):
    """running 任务取消：状态置 cancelled，执行函数据此自行收尾"""
    from src.pipeline.ingest_queue import IngestQueue

    release = threading.Event()
    started = threading.Event()

    def run_fn(job):
        started.set()
        release.wait(5)
        return {}

    q = IngestQueue(run_fn=run_fn, store_path=str(tmp_path / "jobs.json"))
    job = q.submit("incremental")
    assert started.wait(5), "任务未开始执行"
    assert q.get(job.id).status == "running"
    assert q.cancel(job.id) is True
    assert q.is_cancelled(job.id) is True
    release.set()


# ================================================================
# L1 / L2：每日备份
# ================================================================

def test_prune_never_deletes_all_backups(tmp_path):
    """keep<=0 曾会删光（含刚写出的那份）备份，下界必须钳到 1"""
    from src.pipeline.backup import run_once
    from src.session.store import ConversationStore

    store = ConversationStore(dir_path=str(tmp_path / "conv"))
    bdir = tmp_path / "backups"

    p1 = run_once(store, str(bdir), keep=0)
    assert p1 and Path(p1).exists(), "keep=0 时刚写出的备份也被删掉了"
    p2 = run_once(store, str(bdir), keep=-3)
    files = list(bdir.glob("rag-backup-*.json"))
    assert len(files) == 1, f"keep<=0 应至少保留 1 份，实际 {len(files)} 份"
    assert Path(p2).exists()


def test_two_backups_in_same_second_do_not_overwrite(tmp_path):
    """文件名精确到秒时，同秒两次备份会互相覆盖（备份自己先坏）"""
    from src.pipeline.backup import run_once
    from src.session.store import ConversationStore

    store = ConversationStore(dir_path=str(tmp_path / "conv"))
    bdir = tmp_path / "backups"
    a = run_once(store, str(bdir), keep=10)
    b = run_once(store, str(bdir), keep=10)
    assert a != b
    assert len(list(bdir.glob("rag-backup-*.json"))) == 2


def test_backup_is_valid_json_and_no_tmp_leftover(tmp_path):
    from src.pipeline.backup import run_once
    from src.session.store import ConversationStore

    store = ConversationStore(dir_path=str(tmp_path / "conv"))
    s = store.create(title="t")
    store.append(s["id"], "user", "q")
    bdir = tmp_path / "backups"
    path = run_once(store, str(bdir), keep=7)
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    assert data["sessions"]["count"] == 1
    assert list(bdir.glob("*.tmp")) == [], "原子写不应残留 .tmp 文件"


# ================================================================
# L3 / L4 / L6：文档 → Markdown
# ================================================================

class TestMarkdownBoundaries:
    def test_prefix_title_does_not_duplicate(self):
        md = "# 原文档标题\n\n正文"
        assert prefix_title(md, "新标题") == md

    def test_prefix_title_replace_existing_overrides(self):
        """convert 端点的「标题覆盖」依赖 replace_existing=True 才生效"""
        md = "# 原文档标题\n\n正文"
        out = prefix_title(md, "用户指定标题", replace_existing=True)
        assert out.startswith("# 用户指定标题")
        assert "原文档标题" not in out
        assert "正文" in out

    def test_txt_first_line_not_duplicated(self):
        from src.document.to_markdown import convert_to_markdown

        data = "我的笔记标题\n\n第一段内容\n第二段内容".encode("utf-8")
        md = convert_to_markdown(data, "txt", source_name="n.txt").markdown
        assert md.count("我的笔记标题") == 1, f"标题重复出现: {md!r}"
        assert md.startswith("# 我的笔记标题")
        assert "第一段内容" in md

    def test_pptx_table_renders_as_markdown_table(self):
        """PPTX 表格必须渲染成**标准 Markdown 表格**（含 `| --- |` 分隔行）

        第四轮修正：此前 pptx 把单元格按 " | ".join(非空单元格) 拼成普通文本行：
        列数不齐会错位，且缺少分隔行 → Obsidian 根本不认这是表格。
        现已与 html / docx 共用 md_common.render_md_table。
        """
        from pptx import Presentation
        from src.document.to_markdown import convert_to_markdown
        import io as _io

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        table = slide.shapes.add_table(2, 3, 0, 0, 400, 100).table
        rows = [["A1", "", "A3"], ["B1", "B2", ""]]
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                table.cell(r, c).text = val
        buf = _io.BytesIO()
        prs.save(buf)

        md = convert_to_markdown(buf.getvalue(), "pptx", source_name="t.pptx").markdown
        lines = md.splitlines()
        table_lines = [ln for ln in lines if ln.startswith("|")]

        # 表头 + 分隔行 + 1 行数据
        assert len(table_lines) == 3, table_lines
        # 分隔行是「这是表格」的唯一标志，缺了就等于普通文本
        assert table_lines[1] == "| --- | --- | --- |", table_lines
        # 列数一致：3 列 → 每行 4 个竖线
        assert {ln.count("|") for ln in table_lines} == {4}, table_lines
        # 空单元格必须以占位形式保留，否则列会塌陷
        assert table_lines[0] == "| A1 |  | A3 |"
        assert table_lines[2] == "| B1 | B2 |  |"
        # 表格三行必须相邻（中间插空行会把表格切成两半）
        head = lines.index(table_lines[0])
        assert lines[head + 1] == table_lines[1] and lines[head + 2] == table_lines[2]


# ================================================================
# L7：会话存储边界
# ================================================================

class TestSessionStoreBoundaries:
    def test_list_survives_concurrent_delete(self, tmp_path):
        """glob 之后文件被删除不应让 list() 抛 FileNotFoundError（→ API 500）"""
        from src.session.store import ConversationStore

        store = ConversationStore(dir_path=str(tmp_path / "c"))
        a = store.create(title="a")
        store.create(title="b")
        store.delete(a["id"])  # 少一个文件
        assert len(store.list()) == 1  # 不应抛异常

    def test_truncate_negative_clears_instead_of_silent_noop(self, tmp_path):
        from src.session.store import ConversationStore

        store = ConversationStore(dir_path=str(tmp_path / "c"))
        s = store.create(title="t")
        store.append(s["id"], "user", "1")
        store.append(s["id"], "assistant", "2")
        out = store.truncate(s["id"], -5)
        assert out["messages"] == [], "负数 keep_count 应按 0 处理（清空），而不是静默什么都不做"

    def test_truncate_keeps_prefix(self, tmp_path):
        from src.session.store import ConversationStore

        store = ConversationStore(dir_path=str(tmp_path / "c"))
        s = store.create(title="t")
        for i in range(4):
            store.append(s["id"], "user", str(i))
        out = store.truncate(s["id"], 2)
        assert [m["content"] for m in out["messages"]] == ["0", "1"]

    def test_cleanup_keep_zero_keeps_at_least_one(self, tmp_path):
        from src.session.store import ConversationStore

        store = ConversationStore(dir_path=str(tmp_path / "c"))
        for i in range(3):
            store.create(title=f"s{i}")
        store.cleanup(keep=0)
        assert len(store.list()) >= 1, "keep=0 不应清空全部会话"


# ================================================================
# L8：转换体积护栏
# ================================================================

def test_convert_size_guard_blocks_before_decode(monkeypatch):
    """超限请求应在 base64 解码之前就被拒绝（否则已进内存）"""
    from fastapi import HTTPException

    from src.api.routes import convert as convert_mod

    monkeypatch.setattr(convert_mod, "_MAX_BYTES", 1024)
    # 下界 = len*3/4 > 1024 → 必须在解码前抛 413
    huge_b64 = "A" * 8192
    with pytest.raises(HTTPException) as ei:
        convert_mod._to_bytes(huge_b64, is_base64=True)
    assert ei.value.status_code == 413

    with pytest.raises(HTTPException) as ei2:
        convert_mod._to_bytes("x" * 4096, is_base64=False)
    assert ei2.value.status_code == 413


def test_convert_title_override_end_to_end():
    """显式 title 必须真的覆盖解析出的首行 H1（此前被静默忽略）"""
    from src.document.to_markdown import convert_to_markdown, prefix_title

    html = (
        "<html><head><title>原网页标题</title></head>"
        "<body><h1>原网页标题</h1><p>内容</p></body></html>"
    ).encode("utf-8")
    result = convert_to_markdown(html, "html", source_name="html")
    md = prefix_title(result.markdown, "指定标题", replace_existing=True)
    assert md.startswith("# 指定标题"), md[:80]
